from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r"C:\Users\coolf\OneDrive\Documents\claude\portoflio\Acting Resume - Farhan.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_line(cell, title, detail):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    r.bold = True
    p.add_run("\n" + detail)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9)
styles["Normal"].paragraph_format.space_after = Pt(3)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(21)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = None
styles["Heading 2"].font.name = "Aptos Display"
styles["Heading 2"].font.size = Pt(10)
styles["Heading 2"].font.bold = True
styles["Heading 2"].paragraph_format.space_before = Pt(12)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Farhan Naushad")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ACTOR")
r.bold = True
r.font.size = Pt(10)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("farhannaushad020303@gmail.com  |  +1 (613) 484-8348  |  farhannaushad.com")

table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(4.6)
left, right = table.rows[0].cells
left.width = Inches(2.2)
right.width = Inches(4.6)
left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

left.paragraphs[0].text = "DETAILS"
left.paragraphs[0].style = "Heading 2"
for item in (
    "Height  5'9\"",
    "Weight  185 lb",
    "Hair  Black",
    "Eyes  Dark brown",
    "Ethnicity  South Asian",
    "Current facial hair  Beard",
):
    left.add_paragraph(item)

left.add_paragraph("TRAINING", style="Heading 2")
left.add_paragraph("LAMDA Grades 3, 6, 7 and 8 (2016–2020)")

left.add_paragraph("SKILLS", style="Heading 2")
left.add_paragraph("Karate, dancing, swimming, filmmaking, editing, colour grading, sound design, voice acting")

left.add_paragraph("LANGUAGES", style="Heading 2")
left.add_paragraph("English, Malayalam, Hindi; conversational Tamil")

right.paragraphs[0].text = "FILM AND SCREEN"
right.paragraphs[0].style = "Heading 2"
add_line(right, "The Performance of Being Human", "Short film  |  Supporting  |  2026")
add_line(right, "Scarberia", "Short film  |  Supporting  |  2026")
add_line(right, "Alexa", "Short film  |  Lead  |  2025  |  Self-produced")
add_line(right, "The Shave", "Short film  |  Lead  |  2025  |  Self-produced")

right.add_paragraph("THEATRE", style="Heading 2")
add_line(right, "Class of Rock", "High school play  |  Supporting  |  2018")
add_line(right, "I Am Zara Amaya", "High school play  |  Supporting  |  2019")

doc.save(OUT)
print(OUT)
